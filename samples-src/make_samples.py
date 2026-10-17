"""產生範例文件：公版規範（DOCX）與送審規格書（DOCX，再由 Word 轉 PDF）。

內容全為虛構。answer_key.json 記錄預埋的差異，網頁用它計算範例案例的命中率。
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).resolve().parent.parent / "public" / "samples"
OUT.mkdir(parents=True, exist_ok=True)


def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Microsoft JhengHei"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    st.font.size = Pt(11)
    return doc


def title(doc, text, sub=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(18)
    if sub:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(sub).font.size = Pt(10)


def h(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)


def item(doc, text):
    doc.add_paragraph(text)


TEMPLATE = [
    ("1 適用範圍", [
        "1.1 本規範適用於廠區新設或汰換之空氣壓縮機系統，包含空壓機本體、乾燥機、過濾器及監控設備。",
    ]),
    ("2 一般要求", [
        "2.1 設備之設計、製造及安裝應符合 CNS 及 IEC 相關標準。",
        "2.2 保固期間自驗收合格日起算不得少於 24 個月，保固期間內零件及工資均由廠商負擔。",
        "2.3 交貨時應提供中文操作手冊、維護手冊及竣工圖各一式三份。",
        "2.4 廠商應於驗收前提供現場操作及保養教育訓練，時數不少於 8 小時。",
    ]),
    ("3 空壓機本體", [
        "3.1 型式為油冷式螺旋空壓機。",
        "3.2 額定排氣壓力為 0.7 MPa(G)，壓力變動範圍 ±0.02 MPa。",
        "3.3 於額定排氣壓力下，實際排氣量（FAD）不得低於 10 m³/min。",
        "3.4 驅動馬達效率等級應達 IE4。",
        "3.5 應採變頻控制，排氣量可於 40%～100% 範圍內連續調整。",
        "3.6 運轉噪音於機體 1 m 處量測不得超過 72 dB(A)。",
        "3.7 冷卻方式為氣冷式，環境溫度 45°C 以下應能連續運轉。",
    ]),
    ("4 乾燥機與過濾", [
        "4.1 冷凍式乾燥機之壓力露點應達 3°C 以下。",
        "4.2 精密過濾器過濾精度 0.01 μm，出口含油量不得超過 0.01 mg/m³。",
        "4.3 排水器應採零耗氣式自動排水。",
    ]),
    ("5 監控與通訊", [
        "5.1 控制器應配備彩色觸控螢幕，可顯示壓力、溫度、運轉時數及警報。",
        "5.2 控制器應支援 Modbus TCP，並可連線至廠區監控系統（SCADA）。",
        "5.3 應具備高溫、過載及保養到期警報功能。",
        "5.4 運轉資料應可於控制器內保存 12 個月以上。",
    ]),
    ("6 安全", [
        "6.1 安全閥應經合格檢驗並附檢驗證明。",
        "6.2 機體正面應設置緊急停止按鈕。",
        "6.3 壓力容器應取得主管機關檢查合格證明。",
    ]),
    ("7 測試與驗收", [
        "7.1 出廠前應完成工廠測試（FAT）並提供測試報告。",
        "7.2 安裝完成後應進行連續 72 小時運轉測試，期間不得發生異常停機。",
        "7.3 噪音值應於現場量測並記錄於驗收報告。",
    ]),
]


def make_template():
    doc = new_doc()
    title(doc, "空氣壓縮機系統標準規範（公版）", "文件編號 STD-AC-001　版次 1.0　（範例文件，內容為虛構）")
    for sec, items in TEMPLATE:
        h(doc, sec)
        for t in items:
            item(doc, t)
    doc.save(OUT / "公版_空壓機系統標準規範.docx")


def make_submission():
    doc = new_doc()
    title(doc, "空壓機系統規格書（送審版）", "範例機電股份有限公司　提案編號 P-2026-031　（範例文件，內容為虛構）")

    h(doc, "1 產品概述")
    item(doc, "1.1 本案提供油冷式螺旋空壓機一套，含冷凍式乾燥機、精密過濾器及控制系統，設計及製造符合 CNS 及 IEC 相關標準。")
    item(doc, "1.2 控制器支援 Modbus TCP 通訊，可直接連線至廠區監控系統（SCADA）。")

    h(doc, "2 主要規格")
    rows = [
        ("項目", "規格"),
        ("額定排氣壓力", "0.8 MPa(G)，壓力變動範圍 ±0.02 MPa"),
        ("實際排氣量（FAD）", "9.5 m³/min（於額定排氣壓力下）"),
        ("馬達效率等級", "IE3"),
        ("容量調整", "變頻控制，50%～100% 連續調整"),
        ("噪音", "75 dB(A) 以下（距機體 1 m）"),
        ("冷卻方式", "氣冷式，環境溫度 45°C 以下可連續運轉"),
        ("乾燥機壓力露點", "3°C 以下"),
        ("精密過濾器", "過濾精度 0.01 μm，出口含油量 0.01 ppm 以下"),
        ("排水器", "零耗氣式自動排水"),
    ]
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (a, b) in enumerate(rows):
        t.cell(i, 0).text = a
        t.cell(i, 1).text = b
        if i == 0:
            for c in t.rows[0].cells:
                c.paragraphs[0].runs[0].bold = True

    h(doc, "3 控制與監控")
    item(doc, "3.1 控制器配備彩色觸控螢幕，顯示壓力、溫度、運轉時數及警報。")
    item(doc, "3.2 具備高溫、過載及保養到期警報功能。")
    item(doc, "3.3 提供 4G 雲端監控模組及手機 App，可遠端查看運轉狀態。")

    h(doc, "4 安全裝置")
    item(doc, "4.1 安全閥經合格檢驗並附檢驗證明。")
    item(doc, "4.2 壓力容器取得主管機關檢查合格證明。")

    h(doc, "5 節能選配")
    item(doc, "5.1 加裝熱回收裝置，回收壓縮熱供應廠區熱水，熱回收率 70% 以上。")

    h(doc, "6 測試、文件與保固")
    item(doc, "6.1 出廠前完成工廠測試（FAT）並提供測試報告。")
    item(doc, "6.2 安裝完成後進行連續 24 小時運轉測試。")
    item(doc, "6.3 噪音值於現場量測並記錄於驗收報告。")
    item(doc, "6.4 交貨時提供中文操作手冊、維護手冊及竣工圖各一式三份。")
    item(doc, "6.5 保固期間自驗收合格日起算 12 個月，保固期間內零件及工資由本公司負擔。")

    path = OUT / "送審_空壓機系統規格書.docx"
    doc.save(path)
    return path


# 預埋答案：template 欄為公版條號；keyword 用來比對模型輸出的公版條號或原文
ANSWER_KEY = {
    "diff": [
        {"template": "2.2", "label": "保固 24 個月 → 12 個月"},
        {"template": "3.2", "label": "排氣壓力 0.7 → 0.8 MPa"},
        {"template": "3.3", "label": "排氣量 ≥10 → 9.5 m³/min"},
        {"template": "3.4", "label": "馬達 IE4 → IE3"},
        {"template": "3.5", "label": "變頻範圍 40% → 50% 起"},
        {"template": "3.6", "label": "噪音 72 → 75 dB(A)"},
        {"template": "4.2", "label": "含油量單位 mg/m³ → ppm"},
        {"template": "7.2", "label": "運轉測試 72 → 24 小時，且未要求不得異常停機"},
    ],
    "missing": [
        {"template": "2.4", "label": "教育訓練 8 小時"},
        {"template": "5.4", "label": "運轉資料保存 12 個月"},
        {"template": "6.2", "label": "緊急停止按鈕"},
    ],
    "added": [
        {"submission": "3.3", "keyword": "4G", "label": "4G 雲端監控與手機 App"},
        {"submission": "5.1", "keyword": "熱回收", "label": "熱回收裝置"},
    ],
    "traps": [
        {"template": "5.2", "label": "Modbus TCP 移到送審 1.2，不應列為缺漏"},
        {"template": "4.1", "label": "壓力露點相同，不應列為差異"},
        {"template": "2.1", "label": "CNS/IEC 併入送審 1.1，不應列為缺漏"},
    ],
}

if __name__ == "__main__":
    make_template()
    p = make_submission()
    (OUT.parent.parent / "lib").mkdir(exist_ok=True)
    (OUT.parent.parent / "lib" / "answer_key.json").write_text(
        json.dumps(ANSWER_KEY, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("ok", p)
